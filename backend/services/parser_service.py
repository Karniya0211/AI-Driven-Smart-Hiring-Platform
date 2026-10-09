import io
from typing import Optional

def extract_text_from_pdf(file_stream: io.BytesIO) -> str:
    """Extract text content from a PDF file."""
    try:
        from pypdf import PdfReader
        reader = PdfReader(file_stream)
        text_parts = []
        for page in reader.pages:
            extracted = page.extract_text()
            if extracted:
                text_parts.append(extracted)
        return "\n".join(text_parts)
    except Exception as e:
        print(f"Error reading PDF: {e}")
        return ""

def extract_text_from_docx(file_stream: io.BytesIO) -> str:
    """Extract text content from a DOCX file."""
    try:
        import docx
        doc = docx.Document(file_stream)
        text_parts = [para.text for para in doc.paragraphs if para.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text_parts.append(row_text)
        return "\n".join(text_parts)
    except Exception as e:
        print(f"Error reading DOCX: {e}")
        return ""

def parse_resume_file(filename: str, file_bytes: bytes) -> str:
    """Detect file type and extract clean text content."""
    stream = io.BytesIO(file_bytes)
    lower_name = filename.lower()
    if lower_name.endswith(".pdf"):
        text = extract_text_from_pdf(stream)
    elif lower_name.endswith(".docx") or lower_name.endswith(".doc"):
        text = extract_text_from_docx(stream)
    elif lower_name.endswith(".txt"):
        text = file_bytes.decode("utf-8", errors="ignore")
    else:
        # Fallback attempt text decode
        text = file_bytes.decode("utf-8", errors="ignore")
    return text.strip()
